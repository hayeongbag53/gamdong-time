"""Render a saved ProposalHistory row as a downloadable .docx proposal
document, per the user's request that saving a proposal should also
produce a document the admin can hand off / download."""

import json
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

KOREAN_FONT = "맑은 고딕"
ACCENT = RGBColor(0xE8, 0x5A, 0x0A)


def _set_korean_font(run, size=None, bold=None, color=None):
    run.font.name = KOREAN_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), KOREAN_FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = color


def _heading(doc, text, size=15):
    p = doc.add_paragraph()
    p.space_before = Pt(14)
    run = p.add_run(text)
    _set_korean_font(run, size=size, bold=True, color=ACCENT)
    return p


def _body(doc, text, size=10.5):
    p = doc.add_paragraph()
    for i, line in enumerate(str(text or "").split("\n")):
        if i > 0:
            p.add_run().add_break()
        run = p.add_run(line)
        _set_korean_font(run, size=size)
    return p


def _label_value_row(table, label, value):
    row = table.add_row().cells
    run = row[0].paragraphs[0].add_run(label)
    _set_korean_font(run, size=10, bold=True)
    run2 = row[1].paragraphs[0].add_run(str(value if value is not None else "-"))
    _set_korean_font(run2, size=10)


def generate_proposal_docx(proposal):
    products = json.loads(proposal.recommended_products_json or "[]")
    references = json.loads(proposal.reference_summary_json or "[]")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = KOREAN_FONT
    style.element.rPr.rFonts.set(qn("w:eastAsia"), KOREAN_FONT)
    style.font.size = Pt(10.5)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("상품 제안서")
    _set_korean_font(title_run, size=22, bold=True)

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle_p.add_run("감동타임 · AI 고객 제안 챗봇")
    _set_korean_font(subtitle_run, size=10, color=RGBColor(0x88, 0x88, 0x88))

    info_table = doc.add_table(rows=0, cols=2)
    info_table.autofit = True
    _label_value_row(info_table, "고객사", proposal.customer_name_text or "미지정")
    _label_value_row(info_table, "작성자", proposal.admin.name if proposal.admin else "-")
    _label_value_row(info_table, "작성일", proposal.created_at.strftime("%Y-%m-%d %H:%M"))
    if proposal.expected_price:
        _label_value_row(info_table, "예상 단가", proposal.expected_price)

    _heading(doc, "1. 고객 요청 내용")
    _body(doc, proposal.request_text)

    if references:
        _heading(doc, "2. 참고한 과거 유사 사례")
        ref_table = doc.add_table(rows=1, cols=5)
        ref_table.style = "Light Grid Accent 1"
        headers = ["고객사", "수량", "1인당 단가", "만족도", "유사도"]
        for cell, h in zip(ref_table.rows[0].cells, headers):
            run = cell.paragraphs[0].add_run(h)
            _set_korean_font(run, size=9.5, bold=True)
        for ref in references:
            cells = ref_table.add_row().cells
            values = [
                ref.get("company_name", "-"),
                f"{ref.get('quantity', '-')}명",
                f"{int(ref.get('price_per_person', 0)):,}원" if ref.get("price_per_person") else "-",
                str(ref.get("satisfaction_score", "-")),
                f"{ref.get('similarity', '-')}%",
            ]
            for cell, v in zip(cells, values):
                run = cell.paragraphs[0].add_run(v)
                _set_korean_font(run, size=9.5)

    if products:
        _heading(doc, "3. 추천 상품 구성")
        for name in products:
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(name)
            _set_korean_font(run, size=10.5)

    _heading(doc, "4. 최종 제안 내용")
    _body(doc, proposal.final_proposal)

    footer_p = doc.add_paragraph()
    footer_p.space_before = Pt(24)
    footer_run = footer_p.add_run(
        f"본 제안서는 사내 주문 히스토리를 근거로 AI가 초안을 생성하고, "
        f"담당자({proposal.admin.name if proposal.admin else '-'})가 검토·수정한 문서입니다."
    )
    _set_korean_font(footer_run, size=8.5, color=RGBColor(0x99, 0x99, 0x99))

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
